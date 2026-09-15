"""Excel export for Manage Products (all price tiers + live pricing formulas)."""
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db.models import Prefetch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ProductPriceTier
from .resources import ProductResource

MONEY_FORMAT = '0.00'
MARGIN_FORMAT = '0.00'
HEADER_FILL = PatternFill(start_color='E5E7EB', end_color='E5E7EB', fill_type='solid')
HEADER_FONT = Font(bold=True)


def _as_float(value):
    if value is None or value == '':
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _margin_percent(base_cost, price):
    """((price - cost) / price) * 100, or None if it cannot be computed."""
    try:
        cost = Decimal(str(base_cost))
        sell = Decimal(str(price))
    except (InvalidOperation, TypeError, ValueError):
        return None
    if sell == 0:
        return None
    return (sell - cost) / sell * Decimal('100')


def live_price_formula(base_ref, margin_ref):
    """Selling price = base_cost / (1 - profit_margin/100)."""
    return (
        f'=IF(OR({base_ref}="",{margin_ref}=""),"",'
        f'IF({margin_ref}>=100,"",ROUND({base_ref}/(1-{margin_ref}/100),2)))'
    )


def live_margin_formula(base_ref, price_ref):
    """Profit margin % = (selling_price - base_cost) / selling_price * 100."""
    return (
        f'=IF(OR({base_ref}="",{price_ref}="",{price_ref}=0),"",'
        f'ROUND(({price_ref}-{base_ref})/{price_ref}*100,2))'
    )


def _col_letter(index_0):
    return get_column_letter(index_0 + 1)


def _tier_headers(max_tiers):
    headers = []
    for i in range(1, max_tiers + 1):
        headers.extend([
            f'tier_{i}_min_qty',
            f'tier_{i}_profit_margin',
            f'tier_{i}_selling_price',
        ])
    return headers


def prefetch_products_for_export(queryset):
    return queryset.prefetch_related(
        'categories',
        'suppliers',
        Prefetch(
            'price_tiers',
            queryset=ProductPriceTier.objects.order_by('min_quantity'),
        ),
    )


def _apply_live_price_and_margin(cell_price, cell_margin, base_ref, price_ref, margin_ref, has_cost):
    """
    Link selling price and profit margin with live Excel formulas.

    Changing profit_margin or base_cost recalculates selling_price.
    Profit_margin is stored as the current % so it can be edited; it is also
    the input to the selling-price formula. A circular pair (margin formula
    referencing price formula) is avoided by keeping margin as the input value.
    """
    margin_val = cell_margin.value
    price_val = cell_price.value
    if has_cost and margin_val not in (None, ''):
        cell_price.value = live_price_formula(base_ref, margin_ref)
        cell_price.number_format = MONEY_FORMAT
        cell_margin.number_format = MARGIN_FORMAT
    elif has_cost and price_val not in (None, ''):
        cell_margin.value = live_margin_formula(base_ref, price_ref)
        cell_margin.number_format = MARGIN_FORMAT
        cell_price.number_format = MONEY_FORMAT
    else:
        if price_val not in (None, ''):
            cell_price.number_format = MONEY_FORMAT
        if margin_val not in (None, ''):
            cell_margin.number_format = MARGIN_FORMAT


def build_products_xlsx(queryset) -> bytes:
    """
    Build a products workbook with every price tier and live pricing formulas.

    Default columns stay compatible with product upload (sku, selling_price,
    profit_margin, …). Extra `tier_N_*` columns are appended for volume prices.

    Live formulas (when a base cost exists):
    - selling_price = ROUND(base_cost / (1 - profit_margin/100), 2)
    - each tier selling price uses the same formula against that tier's margin
    """
    products = list(prefetch_products_for_export(queryset))
    resource = ProductResource()
    dataset = resource.export(products)

    max_tiers = 0
    for product in products:
        max_tiers = max(max_tiers, len(list(product.price_tiers.all())))

    headers = list(dataset.headers)
    headers.extend(_tier_headers(max_tiers))
    index = {name: i for i, name in enumerate(headers)}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    base_i = index['base_cost']
    price_i = index['selling_price']
    margin_i = index['profit_margin']

    for row_idx, (values, product) in enumerate(zip(dataset, products), start=2):
        row_values = list(values) + [None] * (len(headers) - len(values))

        base_cost = product.base_cost
        selling_price = product.selling_price
        live_margin = None
        if base_cost is not None and selling_price not in (None, ''):
            live_margin = _margin_percent(base_cost, selling_price)
        margin_value = live_margin if live_margin is not None else product.profit_margin

        row_values[base_i] = _as_float(base_cost)
        row_values[margin_i] = _as_float(margin_value)
        row_values[price_i] = _as_float(selling_price)

        tiers = list(product.price_tiers.all())
        for t_i, tier in enumerate(tiers, start=1):
            qty_i = index[f'tier_{t_i}_min_qty']
            t_margin_i = index[f'tier_{t_i}_profit_margin']
            t_price_i = index[f'tier_{t_i}_selling_price']
            row_values[qty_i] = tier.min_quantity
            t_margin = _margin_percent(base_cost, tier.price) if base_cost is not None else None
            row_values[t_margin_i] = _as_float(t_margin)
            row_values[t_price_i] = _as_float(tier.price)

        ws.append(row_values)
        excel_row = row_idx
        base_ref = f'{_col_letter(base_i)}{excel_row}'
        has_cost = base_cost is not None

        if has_cost:
            ws.cell(row=excel_row, column=base_i + 1).number_format = MONEY_FORMAT

        _apply_live_price_and_margin(
            ws.cell(row=excel_row, column=price_i + 1),
            ws.cell(row=excel_row, column=margin_i + 1),
            base_ref,
            f'{_col_letter(price_i)}{excel_row}',
            f'{_col_letter(margin_i)}{excel_row}',
            has_cost,
        )

        for t_i, _tier in enumerate(tiers, start=1):
            t_margin_i = index[f'tier_{t_i}_profit_margin']
            t_price_i = index[f'tier_{t_i}_selling_price']
            _apply_live_price_and_margin(
                ws.cell(row=excel_row, column=t_price_i + 1),
                ws.cell(row=excel_row, column=t_margin_i + 1),
                base_ref,
                f'{_col_letter(t_price_i)}{excel_row}',
                f'{_col_letter(t_margin_i)}{excel_row}',
                has_cost,
            )

    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = min(max(len(str(header)) + 2, 12), 40)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
