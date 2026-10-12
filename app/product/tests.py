import json

from django.contrib.auth import get_user_model
from django.test import Client, TestCase

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from django.core.files.uploadedfile import SimpleUploadedFile
from tablib import Dataset

from inventory.models import Supplier
from product.invoice_sku import (
    build_invoice_sku_preview,
    dataset_is_invoice_item_sheet,
    extract_invoice_item_rows,
    normalize_item_code,
)
from product.models import Product, ProductInvoiceNameAlias, ProductPriceTier
from product.views import _build_merge_candidate_groups, _products_are_merge_candidates
from sales.models import Invoice, InvoiceItem


class ProductMergeSuggestionTests(TestCase):
    def setUp(self):
        self.mounjaro_10 = Product.objects.create(
            name='MOUNJARO® KwikPen®, 1 Stk (English Alternative) 10mg'
        )
        self.mounjaro_15 = Product.objects.create(
            name='MOUNJARO® KwikPen®, 1 Stk (English Alternative) 15mg'
        )
        self.master = Product.objects.create(name='HYALGAN (Fidia)')
        self.duplicate = Product.objects.create(name='Hylagan - (Fidia )')

    def test_different_doses_are_not_merge_candidates(self):
        self.assertFalse(
            _products_are_merge_candidates(self.mounjaro_10.name, self.mounjaro_15.name)
        )
        groups = _build_merge_candidate_groups(
            [self.mounjaro_10, self.mounjaro_15],
        )
        self.assertEqual(groups, [])

    def test_spelling_variants_are_merge_candidates(self):
        self.assertTrue(
            _products_are_merge_candidates(self.master.name, self.duplicate.name)
        )
        groups = _build_merge_candidate_groups([self.master, self.duplicate])
        self.assertEqual(len(groups), 1)
        self.assertEqual(len(groups[0]), 2)


class ProductMergeApiTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='mergeapi',
            password='testpass123',
            is_staff=True,
        )
        self.client = Client()
        self.client.force_login(self.user)
        self.master = Product.objects.create(name='HYALGAN (Fidia)')
        self.duplicate = Product.objects.create(name='Hylagan - (Fidia )')
        self.mounjaro_10 = Product.objects.create(
            name='MOUNJARO® KwikPen®, 1 Stk (English Alternative) 10mg'
        )
        self.mounjaro_15 = Product.objects.create(
            name='MOUNJARO® KwikPen®, 1 Stk (English Alternative) 15mg'
        )

    def test_merge_suggestions_exclude_dose_variants(self):
        response = self.client.get(
            '/api/product-merge-suggestions/',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(response.status_code, 200)
        groups = response.json()['groups']
        mounjaro_group = next(
            (group for group in groups if len(group['products']) >= 2
             and any('MOUNJARO' in product['name'] for product in group['products'])),
            None,
        )
        self.assertIsNone(mounjaro_group)

    def test_merge_suggestions_find_spelling_duplicates(self):
        response = self.client.get(
            '/api/product-merge-suggestions/',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(response.status_code, 200)
        hyalgan_groups = [
            group for group in response.json()['groups']
            if {product['id'] for product in group['products']} == {self.master.pk, self.duplicate.pk}
        ]
        self.assertEqual(len(hyalgan_groups), 1)

    def test_merge_suggestions_selection_mode_returns_manual_group(self):
        unrelated_a = Product.objects.create(name='Alpha Widget')
        unrelated_b = Product.objects.create(name='Beta Gadget')
        response = self.client.get(
            f'/api/product-merge-suggestions/?product_ids={unrelated_a.pk},{unrelated_b.pk}',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['mode'], 'selection')
        self.assertEqual(len(data['groups']), 1)
        self.assertEqual(len(data['groups'][0]['products']), 2)


class ProductMergeTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='mergetest',
            password='testpass123',
            is_staff=True,
        )
        self.client = Client()
        self.client.force_login(self.user)
        self.supplier = Supplier.objects.create(name='MBA Pharmaceuticals Pvt Ltd')
        self.master = Product.objects.create(name='HYALGAN (Fidia)')
        self.duplicate = Product.objects.create(name='Hylagan - (Fidia )')
        self.invoice = Invoice.objects.create(
            invoice_id='PI/G-TEST001',
            supplier=self.supplier,
        )
        InvoiceItem.objects.create(
            invoice=self.invoice,
            product=self.duplicate,
            description=self.duplicate.name,
            quantity=200,
            unit_price='90.77',
        )

    def test_merge_reassigns_invoice_items_to_master(self):
        response = self.client.post(
            '/api/merge-products/',
            data=json.dumps({
                'primary_id': self.master.pk,
                'merge_ids': [self.duplicate.pk],
            }),
            content_type='application/json',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])
        self.assertFalse(Product.objects.filter(pk=self.duplicate.pk).exists())
        item = InvoiceItem.objects.get(pk=self.invoice.items.first().pk)
        self.assertEqual(item.product_id, self.master.pk)
        from product.models import ProductInvoiceNameAlias
        alias = ProductInvoiceNameAlias.objects.get(product=self.master)
        self.assertEqual(alias.name, 'Hylagan - (Fidia )')

    def test_merge_keeps_nested_invoice_name_aliases(self):
        from product.models import ProductInvoiceNameAlias
        from product.name_aliases import upsert_invoice_name_alias
        older = Product.objects.create(name='Older invoice name')
        upsert_invoice_name_alias(self.duplicate, 'PI line: Hyalgan inj')
        response = self.client.post(
            '/api/merge-products/',
            data=json.dumps({
                'primary_id': self.master.pk,
                'merge_ids': [self.duplicate.pk, older.pk],
            }),
            content_type='application/json',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(response.status_code, 200)
        names = set(
            ProductInvoiceNameAlias.objects.filter(product=self.master).values_list('name', flat=True)
        )
        self.assertIn('Hylagan - (Fidia )', names)
        self.assertIn('Older invoice name', names)
        self.assertIn('PI line: Hyalgan inj', names)


class ProductExcelExportTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='exportstaff',
            password='testpass123',
            is_staff=True,
        )
        self.client = Client()
        self.client.force_login(self.user)
        self.supplier = Supplier.objects.create(name='Export Test Supplier')
        self.product = Product.objects.create(
            name='Tiered Export Product',
            sku='EXP-TIER-1',
            selling_price=Decimal('100.00'),
            profit_margin=Decimal('20.00'),
            saved_base_cost=Decimal('80.00'),
            saved_base_cost_supplier=self.supplier,
        )
        ProductPriceTier.objects.create(
            product=self.product, min_quantity=10, price=Decimal('95.00')
        )
        ProductPriceTier.objects.create(
            product=self.product, min_quantity=20, price=Decimal('90.00')
        )
        self.plain = Product.objects.create(
            name='No Tiers Product',
            sku='EXP-PLAIN-1',
            selling_price=Decimal('50.00'),
            saved_base_cost=Decimal('40.00'),
            saved_base_cost_supplier=self.supplier,
        )

    def _workbook(self, products=None):
        from product.excel_export import build_products_xlsx

        qs = Product.objects.filter(
            pk__in=[p.pk for p in (products or [self.product, self.plain])]
        ).order_by('name')
        return load_workbook(BytesIO(build_products_xlsx(qs)))

    def test_export_includes_all_price_tiers(self):
        wb = self._workbook()
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn('tier_1_min_qty', headers)
        self.assertIn('tier_1_profit_margin', headers)
        self.assertIn('tier_1_selling_price', headers)
        self.assertIn('tier_2_min_qty', headers)
        self.assertIn('tier_2_selling_price', headers)

        rows = {
            ws.cell(row=r, column=headers.index('sku') + 1).value: r
            for r in range(2, ws.max_row + 1)
        }
        tiered_row = rows['EXP-TIER-1']
        self.assertEqual(
            ws.cell(row=tiered_row, column=headers.index('tier_1_min_qty') + 1).value, 10
        )
        self.assertEqual(
            ws.cell(row=tiered_row, column=headers.index('tier_2_min_qty') + 1).value, 20
        )

        plain_row = rows['EXP-PLAIN-1']
        self.assertIsNone(
            ws.cell(row=plain_row, column=headers.index('tier_1_min_qty') + 1).value
        )

    def test_export_uses_live_formulas_for_price_and_margin(self):
        wb = self._workbook([self.product])
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        row = 2
        base_col = headers.index('base_cost') + 1
        price_col = headers.index('selling_price') + 1
        margin_col = headers.index('profit_margin') + 1

        base_ref = f'{get_column_letter(base_col)}{row}'
        margin_ref = f'{get_column_letter(margin_col)}{row}'
        price_cell = ws.cell(row=row, column=price_col)
        self.assertIsInstance(price_cell.value, str)
        self.assertTrue(price_cell.value.startswith('='))
        self.assertIn(base_ref, price_cell.value)
        self.assertIn(margin_ref, price_cell.value)
        self.assertIn('/(1-', price_cell.value.replace(' ', ''))

        t1_price_col = headers.index('tier_1_selling_price') + 1
        t1_margin_col = headers.index('tier_1_profit_margin') + 1
        t1_price = ws.cell(row=row, column=t1_price_col)
        self.assertIsInstance(t1_price.value, str)
        self.assertTrue(t1_price.value.startswith('='))
        self.assertIn(base_ref, t1_price.value)
        self.assertIn(f'{get_column_letter(t1_margin_col)}{row}', t1_price.value)

        t2_price = ws.cell(row=row, column=headers.index('tier_2_selling_price') + 1)
        self.assertIsInstance(t2_price.value, str)
        self.assertTrue(t2_price.value.startswith('='))

    def test_export_selected_products_endpoint(self):
        response = self.client.get(f'/export-products/?ids={self.product.pk}')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        sku_col = headers.index('sku') + 1
        skus = [ws.cell(row=r, column=sku_col).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(skus, ['EXP-TIER-1'])
        self.assertIn('tier_1_min_qty', headers)
        price_col = headers.index('selling_price') + 1
        self.assertTrue(str(ws.cell(row=2, column=price_col).value).startswith('='))


class ProductImportFormulaTests(TestCase):
    def test_formula_cells_resolve_from_margin_and_cost(self):
        from product.resources import ProductResource

        resource = ProductResource()
        row = {
            'base_cost': '80',
            'profit_margin': '20',
            'selling_price': '=IF(OR(J2="",L2=""),"",IF(L2>=100,"",ROUND(J2/(1-L2/100),2)))',
            'tier_1_min_qty': 10,
            'tier_1_profit_margin': '15.79',
            'tier_1_selling_price': '=ROUND(J2/(1-N2/100),2)',
        }
        resource._resolve_pricing_formula_cells(row)
        self.assertEqual(row['selling_price'], Decimal('100.00'))
        self.assertEqual(row['tier_1_selling_price'], Decimal('95.00'))

    def test_blank_tier_columns_do_not_clear_existing_tiers(self):
        from product.resources import ProductResource

        resource = ProductResource()
        self.assertIsNone(resource._parse_price_tiers_from_row({
            'sku': 'X',
            'tier_1_min_qty': '',
            'tier_1_selling_price': '',
            'tier_1_profit_margin': '',
        }))


class InvoiceSkuUploadTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='skuupload',
            password='testpass123',
            is_staff=True,
        )
        self.client = Client()
        self.client.force_login(self.user)
        self.product = Product.objects.create(name='Re N Tox 100u')
        self.other = Product.objects.create(name='REJURAN Healer', sku='OLD-SKU')

    def _sheet(self, headers, data_rows):
        dataset = Dataset()
        dataset.headers = headers
        for row in data_rows:
            dataset.append(row)
        return dataset

    def _xlsx_bytes(self, headers, data_rows):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in data_rows:
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_code_description_price_sheet_imports_reference_price(self):
        dataset = self._sheet(
            ['Code', 'Description', 'Price'],
            [
                [89693, 'Re N Tox 100U', 109],
                [89693, 'Re N Tox 100U', 109],
                [89639, 'Innotox 50U', 68],
            ],
        )
        self.assertTrue(dataset_is_invoice_item_sheet(dataset))
        rows = extract_invoice_item_rows(dataset, require_item_code=False)
        by_desc = {row['description']: row for row in rows}
        self.assertEqual(len(by_desc), 2)
        self.assertEqual(by_desc['Re N Tox 100U']['item_code'], '89693')
        self.assertEqual(by_desc['Re N Tox 100U']['reference_price'], '109.00')
        self.assertEqual(by_desc['Innotox 50U']['reference_price'], '68.00')

        upload = SimpleUploadedFile(
            'pd-local-sales-mapping.xlsx',
            self._xlsx_bytes(['Code', 'Description', 'Price'], [[89693, 'Re N Tox 100U', 109]]),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        preview = self.client.post(
            '/api/product-alias-import/',
            {'file': upload},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        row = preview.json()['rows'][0]
        self.assertEqual(row['reference_price'], '109.00')
        self.assertEqual(row['product_id'], self.product.pk)

    def test_item_code_sheet_is_detected_and_numeric_codes_are_clean(self):
        titled = Dataset()
        titled.headers = ['Payable Invoice Detail', '']
        titled.append(['Item Code', 'Description'])
        titled.append([89693.0, 'Re N Tox 100u'])
        self.assertTrue(dataset_is_invoice_item_sheet(titled))
        rows = extract_invoice_item_rows(titled)
        self.assertEqual(rows[0]['item_code'], '89693')
        self.assertEqual(rows[0]['description'], 'Re N Tox 100u')
        self.assertEqual(normalize_item_code(89694), '89694')

        catalog = self._sheet(['sku', 'name', 'description'], [['A', 'B', 'C']])
        self.assertFalse(dataset_is_invoice_item_sheet(catalog))

    def test_preview_matches_name_alias_and_conflict(self):
        ProductInvoiceNameAlias.objects.create(
            product=self.other,
            name='REJURAN® Healer',
            name_normalized=ProductInvoiceNameAlias.normalize_name('REJURAN® Healer'),
        )
        taken = Product.objects.create(name='Already coded', sku='100017')
        blocked = Product.objects.create(name='Bellast Ultra')
        dataset = self._sheet(
            ['Item Code', 'Description'],
            [
                ['89693', 'Re N Tox 100u'],
                ['89710', 'REJURAN® Healer'],
                ['100017', 'Bellast Ultra'],
                ['', 'Special Discount'],
            ],
        )
        preview = {row['description']: row for row in build_invoice_sku_preview(extract_invoice_item_rows(dataset))}
        self.assertEqual(preview['Re N Tox 100u']['status'], 'update')
        self.assertEqual(preview['Re N Tox 100u']['product_id'], self.product.pk)
        self.assertEqual(preview['REJURAN® Healer']['status'], 'update')
        self.assertEqual(preview['REJURAN® Healer']['product_id'], self.other.pk)
        self.assertEqual(preview['REJURAN® Healer']['current_sku'], 'OLD-SKU')
        self.assertEqual(preview['Bellast Ultra']['status'], 'conflict')
        self.assertIn('Already coded', preview['Bellast Ultra']['message'])
        self.assertNotIn('Special Discount', preview)
        self.assertEqual(blocked.name, 'Bellast Ultra')
        self.assertEqual(taken.sku, '100017')

    def test_upload_sets_sku_and_saves_mapping_for_unmatched_choice(self):
        upload = SimpleUploadedFile(
            'items.csv',
            b'Item Code,Description\n89693,Re N Tox 100u\n89710,REJURAN(r) Healer\n',
            content_type='text/csv',
        )
        preview = self.client.post(
            '/upload-products/',
            {'file': upload},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        data = preview.json()
        self.assertEqual(data['mode'], 'invoice_sku')
        by_desc = {row['description']: row for row in data['preview_data']}
        self.assertEqual(by_desc['Re N Tox 100u']['status'], 'update')
        self.assertEqual(by_desc['REJURAN(r) Healer']['status'], 'unmatched')

        confirm = self.client.post(
            '/upload-products/',
            {
                'confirm': 'true',
                'mode': 'invoice_sku',
                'temp_file_path': data['temp_file_path'],
                'assignments': json.dumps([
                    {'row_key': by_desc['Re N Tox 100u']['row_key'], 'product_id': self.product.pk},
                    {'row_key': by_desc['REJURAN(r) Healer']['row_key'], 'product_id': self.other.pk},
                ]),
            },
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(confirm.status_code, 200, confirm.content)
        self.product.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual(self.product.sku, '89693')
        self.assertEqual(self.other.sku, '89710')
        alias = ProductInvoiceNameAlias.objects.get(
            name_normalized=ProductInvoiceNameAlias.normalize_name('REJURAN(r) Healer')
        )
        self.assertEqual(alias.product_id, self.other.pk)
        self.assertFalse(ProductInvoiceNameAlias.objects.filter(
            name_normalized=ProductInvoiceNameAlias.normalize_name('Re N Tox 100u')
        ).exists())

    def test_mapping_api_creates_lists_and_deletes_alias(self):
        created = self.client.post(
            '/api/product-aliases/',
            data=json.dumps({'name': 'Xeomin 100u', 'product_id': self.product.pk}),
            content_type='application/json',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(created.status_code, 200, created.content)
        alias_id = created.json()['item']['id']
        self.assertEqual(created.json()['item']['source'], 'manual')

        listed = self.client.get(
            '/api/product-aliases/?search=Xeomin',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()['pagination']['total_count'], 1)

        deleted = self.client.post(
            f'/api/product-aliases/{alias_id}/',
            data=json.dumps({'delete': True}),
            content_type='application/json',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(deleted.status_code, 200)
        self.assertFalse(ProductInvoiceNameAlias.objects.filter(pk=alias_id).exists())

    def test_mapping_import_reviews_descriptions_without_changing_sku(self):
        from product.invoice_sku import build_invoice_mapping_preview

        self.other.alias_name = 'Xeomin vial'
        self.other.save(update_fields=['alias_name'])
        ProductInvoiceNameAlias.objects.create(
            product=self.product,
            name='Re N Tox 100 units',
            name_normalized=ProductInvoiceNameAlias.normalize_name('Re N Tox 100 units'),
            source=ProductInvoiceNameAlias.SOURCE_MANUAL,
        )
        dataset = self._sheet(
            ['Item Code', 'Description'],
            [
                ['89693', 'Re N Tox 100u'],
                ['1', 'Re N Tox 100 units'],
                ['2', 'Xeomin vial'],
                ['', 'Brand new line'],
            ],
        )
        preview = {row['description']: row for row in build_invoice_mapping_preview(
            extract_invoice_item_rows(dataset, require_item_code=False)
        )}
        self.assertEqual(preview['Re N Tox 100u']['status'], 'catalog')
        self.assertEqual(preview['Re N Tox 100 units']['status'], 'mapped')
        self.assertEqual(preview['Xeomin vial']['status'], 'suggested')
        self.assertEqual(preview['Brand new line']['status'], 'unmatched')
        self.assertEqual(preview['Brand new line']['item_code'], '')

        upload = SimpleUploadedFile(
            'map.csv',
            b'Item Code,Description\n89693,Re N Tox 100u\n2,Xeomin vial\n,Brand new line\n',
            content_type='text/csv',
        )
        response = self.client.post(
            '/api/product-alias-import/',
            {'file': upload},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        by_desc = {row['description']: row for row in data['rows']}
        confirm = self.client.post(
            '/api/product-alias-import/',
            {
                'confirm': 'true',
                'temp_file_path': data['temp_file_path'],
                'assignments': json.dumps([
                    {'row_key': by_desc['Xeomin vial']['row_key'], 'product_id': self.other.pk},
                    {'row_key': by_desc['Brand new line']['row_key'], 'product_id': self.product.pk},
                ]),
            },
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(confirm.status_code, 200, confirm.content)
        self.assertEqual(confirm.json()['created'], 2)
        self.other.refresh_from_db()
        self.product.refresh_from_db()
        self.assertEqual(self.other.sku, 'OLD-SKU')
        self.assertIsNone(self.product.sku)
        self.assertTrue(ProductInvoiceNameAlias.objects.filter(
            name_normalized=ProductInvoiceNameAlias.normalize_name('Brand new line'),
            product=self.product,
        ).exists())

    def test_mapping_lists_every_product_and_saves_reference_price(self):
        bare = Product.objects.create(name='Unaliased Serum ZZZ')
        listed = self.client.get(
            '/api/product-aliases/?search=Unaliased Serum',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(listed.status_code, 200, listed.content)
        items = listed.json()['items']
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]['id'], bare.pk)
        self.assertEqual(items[0]['aliases'], [])
        self.assertEqual(items[0]['reference_price'], '')

        saved = self.client.post(
            f'/api/product-reference-price/{bare.pk}/',
            data=json.dumps({'reference_price': '88.5'}),
            content_type='application/json',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(saved.status_code, 200, saved.content)
        self.assertEqual(saved.json()['reference_price'], '88.50')
        bare.refresh_from_db()
        self.assertEqual(bare.reference_price, Decimal('88.50'))

        cleared = self.client.post(
            f'/api/product-reference-price/{bare.pk}/',
            data=json.dumps({'reference_price': ''}),
            content_type='application/json',
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(cleared.status_code, 200, cleared.content)
        bare.refresh_from_db()
        self.assertIsNone(bare.reference_price)

        page = self.client.get('/manage/', HTTP_HOST='localhost')
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'Reference price')
        self.assertContains(page, 'No alias')
        self.assertContains(page, 'Remove alias')
        self.assertContains(page, 'Add alias')
        self.assertContains(page, 'x-show="aliasOpen"')
        self.assertNotContains(page, 'Add mapping')

    def test_mapping_import_saves_reference_price_without_changing_sku(self):
        price_file = SimpleUploadedFile(
            'prices.csv',
            b'SKU,Reference Price\nOLD-SKU,42.00\n',
            content_type='text/csv',
        )
        preview = self.client.post(
            '/api/product-alias-import/',
            {'file': price_file},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        row = preview.json()['rows'][0]
        self.assertEqual(row['product_id'], self.other.pk)
        self.assertEqual(row['reference_price'], '42.00')
        self.assertEqual(row['status'], 'catalog')

        confirm = self.client.post(
            '/api/product-alias-import/',
            {
                'confirm': 'true',
                'temp_file_path': preview.json()['temp_file_path'],
                'assignments': json.dumps([{
                    'row_key': row['row_key'],
                    'product_id': self.other.pk,
                    'reference_price': '42.00',
                }]),
            },
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(confirm.status_code, 200, confirm.content)
        self.assertEqual(confirm.json()['prices'], 1)
        self.assertEqual(confirm.json()['created'], 0)
        self.other.refresh_from_db()
        self.assertEqual(self.other.reference_price, Decimal('42.00'))
        self.assertEqual(self.other.sku, 'OLD-SKU')
        self.assertFalse(ProductInvoiceNameAlias.objects.filter(product=self.other).exists())

        invoice = SimpleUploadedFile(
            'invoice.csv',
            b'Item Code,Description,Unit Price\n89693,Re N Tox 100u,15.5\n',
            content_type='text/csv',
        )
        invoice_preview = self.client.post(
            '/api/product-alias-import/',
            {'file': invoice},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(invoice_preview.status_code, 200, invoice_preview.content)
        invoice_row = invoice_preview.json()['rows'][0]
        self.assertEqual(invoice_row['status'], 'catalog')
        self.assertEqual(invoice_row['reference_price'], '15.50')
        invoice_confirm = self.client.post(
            '/api/product-alias-import/',
            {
                'confirm': 'true',
                'temp_file_path': invoice_preview.json()['temp_file_path'],
                'assignments': json.dumps([{
                    'row_key': invoice_row['row_key'],
                    'product_id': self.product.pk,
                    'reference_price': '15.50',
                }]),
            },
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(invoice_confirm.status_code, 200, invoice_confirm.content)
        self.product.refresh_from_db()
        self.assertEqual(self.product.reference_price, Decimal('15.50'))
        self.assertIsNone(self.product.sku)
        self.assertFalse(ProductInvoiceNameAlias.objects.filter(
            name_normalized=ProductInvoiceNameAlias.normalize_name('Re N Tox 100u')
        ).exists())
