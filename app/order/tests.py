import json
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.utils import timezone
from openpyxl import load_workbook

from order.models import (
    AgentCommissionPaymentEntry,
    CashBankReceiptEntry,
    Order,
    OrderItem,
    RevenueAdjustmentEntry,
)
from product.models import Product


class OrdersBreakdownProfitMarginTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.superuser = user_model.objects.create_superuser(
            username='breakadmin',
            email='breakadmin@example.com',
            password='testpass123',
        )
        self.staff = user_model.objects.create_user(
            username='breakstaff',
            email='breakstaff@example.com',
            password='testpass123',
            is_staff=True,
        )
        self.agent = user_model.objects.create_user(
            username='breakagent',
            email='breakagent@example.com',
            password='testpass123',
        )
        self.product = Product.objects.create(name='Breakdown Widget', sku='BD-W-1')
        now = timezone.now()
        self.order = Order.objects.create(
            agent=self.agent,
            status=Order.OrderStatus.COMPLETED,
            transaction_date=now.date(),
        )
        OrderItem.objects.create(
            order=self.order,
            product=self.product,
            quantity=2,
            selling_price=Decimal('100.00'),
            landed_cost=Decimal('80.00'),
        )
        self.url = '/order/api/orders-breakdown/'
        self.params = {
            'month': now.month,
            'year': now.year,
        }

    def test_superuser_sees_profit_margin(self):
        client = Client()
        client.force_login(self.superuser)
        response = client.get(self.url, self.params)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['total_revenue'], 200.0)
        self.assertEqual(data['total_profit'], 40.0)
        self.assertEqual(data['total_profit_margin'], 20.0)
        self.assertEqual(len(data['breakdown']), 1)
        row = data['breakdown'][0]
        self.assertEqual(row['profit'], 40.0)
        self.assertEqual(row['profit_margin'], 20.0)

    def test_staff_does_not_receive_profit_fields(self):
        client = Client()
        client.force_login(self.staff)
        response = client.get(self.url, self.params)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertNotIn('total_profit', data)
        self.assertNotIn('total_profit_margin', data)
        self.assertNotIn('profit', data['breakdown'][0])
        self.assertNotIn('profit_margin', data['breakdown'][0])


class CashAndPendingFilterScopeTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.superuser = user_model.objects.create_superuser(
            username='cashadmin',
            email='cashadmin@example.com',
            password='testpass123',
        )
        self.agent = user_model.objects.create_user(
            username='cashagent',
            email='cashagent@example.com',
            password='testpass123',
        )
        self.product = Product.objects.create(name='Cash Widget', sku='CASH-W-1')
        today = timezone.localdate()
        self.this_month = date(today.year, today.month, 1)
        if today.month == 1:
            self.other_month = date(today.year - 1, 12, 15)
        else:
            self.other_month = date(today.year, today.month - 1, 15)

        this_order = Order.objects.create(
            agent=self.agent,
            created_by=self.superuser,
            status=Order.OrderStatus.COMPLETED,
            customer_name='In Period Customer',
            transaction_date=self.this_month,
        )
        OrderItem.objects.create(
            order=this_order,
            product=self.product,
            quantity=2,
            selling_price=Decimal('100.00'),
            landed_cost=Decimal('40.00'),
        )
        other_order = Order.objects.create(
            agent=self.agent,
            created_by=self.superuser,
            status=Order.OrderStatus.COMPLETED,
            customer_name='Other Period Customer',
            transaction_date=self.other_month,
        )
        OrderItem.objects.create(
            order=other_order,
            product=self.product,
            quantity=5,
            selling_price=Decimal('100.00'),
            landed_cost=Decimal('40.00'),
        )

        CashBankReceiptEntry.objects.create(
            payment_type=CashBankReceiptEntry.PaymentType.CASH,
            received_from='In Period Customer',
            collected_by=self.agent,
            transaction_date=self.this_month,
            amount=Decimal('50.00'),
            recorded_by=self.superuser,
        )
        CashBankReceiptEntry.objects.create(
            payment_type=CashBankReceiptEntry.PaymentType.BANK,
            received_from='Other Period Customer',
            collected_by=self.agent,
            transaction_date=self.other_month,
            amount=Decimal('80.00'),
            recorded_by=self.superuser,
        )
        AgentCommissionPaymentEntry.objects.create(
            paid_to=self.agent.username,
            payment_date=self.this_month,
            amount=Decimal('10.00'),
            recorded_by=self.superuser,
        )
        AgentCommissionPaymentEntry.objects.create(
            paid_to=self.agent.username,
            payment_date=self.other_month,
            amount=Decimal('30.00'),
            recorded_by=self.superuser,
        )
        RevenueAdjustmentEntry.objects.create(
            adjustment_type=RevenueAdjustmentEntry.AdjustmentType.LOAN_INTEREST,
            reference='In period interest',
            transaction_date=self.this_month,
            amount=Decimal('7.00'),
            recorded_by=self.superuser,
        )
        RevenueAdjustmentEntry.objects.create(
            adjustment_type=RevenueAdjustmentEntry.AdjustmentType.LOAN_INTEREST,
            reference='Other period interest',
            transaction_date=self.other_month,
            amount=Decimal('9.00'),
            recorded_by=self.superuser,
        )
        self.client = Client()
        self.client.force_login(self.superuser)
        self.period = {'month': today.month, 'year': today.year}

    def test_cards_follow_month_filter(self):
        response = self.client.get('/order/api/manage-orders/', self.period)
        self.assertEqual(response.status_code, 200)
        financial = response.json()['financial']
        self.assertEqual(financial['total_cash_received'], 50.0)
        self.assertEqual(financial['total_commission_paid'], 10.0)
        self.assertEqual(financial['pending_cash_receivable'], 140.0)

    def test_all_periods_includes_every_entry(self):
        response = self.client.get('/order/api/manage-orders/', {'month': 0, 'year': 0})
        self.assertEqual(response.status_code, 200)
        financial = response.json()['financial']
        self.assertEqual(financial['total_cash_received'], 130.0)
        self.assertEqual(financial['total_commission_paid'], 40.0)
        self.assertEqual(financial['pending_cash_receivable'], 530.0)

    def test_view_modals_follow_month_filter(self):
        cash = self.client.get('/order/api/cash-received-transactions/', self.period)
        self.assertEqual(cash.status_code, 200)
        cash_data = cash.json()
        self.assertEqual(cash_data['total'], 50.0)
        self.assertEqual(len(cash_data['transactions']), 1)

        commission = self.client.get('/order/api/commission-paid-transactions/', self.period)
        self.assertEqual(commission.status_code, 200)
        commission_data = commission.json()
        self.assertEqual(commission_data['total'], 10.0)
        self.assertEqual(len(commission_data['transactions']), 1)

        breakdown = self.client.get('/order/api/cash-received-breakdown/', self.period)
        self.assertEqual(breakdown.status_code, 200)
        self.assertEqual(breakdown.json()['total'], 50.0)

        adjustments = self.client.get('/order/api/revenue-adjustment-transactions/', self.period)
        self.assertEqual(adjustments.status_code, 200)
        adjustment_data = adjustments.json()
        self.assertEqual(adjustment_data['total'], 7.0)
        self.assertEqual(len(adjustment_data['transactions']), 1)

    def test_date_range_overrides_month(self):
        response = self.client.get('/order/api/manage-orders/', {
            'month': self.this_month.month,
            'year': self.this_month.year,
            'start_date': self.other_month.isoformat(),
            'end_date': self.other_month.isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        financial = response.json()['financial']
        self.assertEqual(financial['total_cash_received'], 80.0)
        self.assertEqual(financial['pending_cash_receivable'], 390.0)

    def test_positive_line_revenue_requires_and_filters_by_agent(self):
        payload = {
            'adjustment_type': 'COMMISSION_RELEASED',
            'reference': 'Release',
            'transaction_date': self.this_month.isoformat(),
            'amount': '15.00',
        }
        missing = self.client.post(
            '/order/api/revenue-adjustment/',
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertEqual(missing.status_code, 400)

        payload['agent_id'] = self.agent.pk
        created = self.client.post(
            '/order/api/revenue-adjustment/',
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertEqual(created.json()['entry']['agent'], self.agent.username)

        other = get_user_model().objects.create_user(
            username='otheragent',
            email='otheragent@example.com',
            password='testpass123',
        )
        RevenueAdjustmentEntry.objects.create(
            adjustment_type=RevenueAdjustmentEntry.AdjustmentType.COMMISSION_RELEASED,
            agent=other,
            reference='Other agent release',
            transaction_date=self.this_month,
            amount=Decimal('4.00'),
            recorded_by=self.superuser,
        )
        listed = self.client.get('/order/api/revenue-adjustment-transactions/', {
            **self.period,
            'agent': self.agent.pk,
        })
        self.assertEqual(listed.status_code, 200)
        data = listed.json()
        self.assertEqual(data['total'], 15.0)
        self.assertEqual([tx['agent'] for tx in data['transactions']], [self.agent.username])


class OrderStatementSalesOnlyExportTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.superuser = user_model.objects.create_superuser(
            username='stmtadmin',
            email='stmtadmin@example.com',
            password='testpass123',
        )
        self.agent = user_model.objects.create_user(
            username='stmtagent',
            email='stmtagent@example.com',
            password='testpass123',
        )
        today = timezone.localdate()
        self.month = today.month
        self.year = today.year
        self.when = date(today.year, today.month, 1)
        product = Product.objects.create(
            name='Sales Widget',
            sku='STMT-W-1',
            reference_price=Decimal('109.00'),
        )
        order = Order.objects.create(
            agent=self.agent,
            created_by=self.superuser,
            status=Order.OrderStatus.COMPLETED,
            customer_name='Statement Customer',
            transaction_date=self.when,
        )
        OrderItem.objects.create(
            order=order,
            product=product,
            quantity=1,
            selling_price=Decimal('40.00'),
            landed_cost=Decimal('10.00'),
        )
        CashBankReceiptEntry.objects.create(
            payment_type=CashBankReceiptEntry.PaymentType.BANK,
            received_from='Statement Customer',
            collected_by=self.agent,
            transaction_date=self.when,
            amount=Decimal('20.00'),
            recorded_by=self.superuser,
        )
        CashBankReceiptEntry.objects.create(
            payment_type=CashBankReceiptEntry.PaymentType.LOAN,
            received_from='Statement Customer',
            collected_by=self.agent,
            transaction_date=self.when,
            amount=Decimal('5.00'),
            recorded_by=self.superuser,
        )
        RevenueAdjustmentEntry.objects.create(
            adjustment_type=RevenueAdjustmentEntry.AdjustmentType.LOAN_INTEREST,
            agent=self.agent,
            reference='Loan interest',
            transaction_date=self.when,
            amount=Decimal('3.00'),
            recorded_by=self.superuser,
        )
        AgentCommissionPaymentEntry.objects.create(
            paid_to=self.agent.username,
            payment_date=self.when,
            amount=Decimal('2.00'),
            recorded_by=self.superuser,
        )
        self.client = Client()
        self.client.force_login(self.superuser)

    def _workbook(self, **params):
        query = {'month': self.month, 'year': self.year}
        query.update(params)
        response = self.client.get('/order/export/', query)
        self.assertEqual(response.status_code, 200, response.content[:200])
        return load_workbook(BytesIO(response.content))

    def test_full_statement_includes_finance_lines(self):
        sheet = self._workbook().active
        self.assertEqual(sheet.title, 'Orders')
        labels = [row[4] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertIn('Sales Widget', labels)
        self.assertIn('Bank transfer', labels)
        self.assertIn('Loan repayment', labels)
        self.assertIn('Interest of loan', labels)
        self.assertIn('Commission paid', labels)
        headers = [cell.value for cell in sheet[1]]
        self.assertIn('Product Base Cost', headers)
        self.assertIn('Accumulated Profit', headers)

    def test_simple_report_uses_reference_price_and_live_formulas(self):
        workbook = self._workbook(simple_report='1', usd_rate='4.70')
        self.assertEqual(
            workbook.sheetnames,
            ['Sale Transactions', 'Monthly Sales', 'Product Sales'],
        )
        sheet = workbook['Sale Transactions']
        self.assertEqual(sheet['A1'].value, 'Convert rate (MYR per 1 USD)')
        self.assertEqual(sheet['B1'].value, 4.7)
        headers = [cell.value for cell in sheet[3]]
        self.assertEqual(headers, [
            'Order ID', 'Order Date', 'Customer Name', 'Product ID', 'Product Name',
            'Price', 'Quantity', 'Line Revenue (MYR)', 'Line Revenue (USD)',
        ])
        self.assertNotIn('Salesteam', headers)
        self.assertNotIn('Product Base Cost', headers)
        self.assertNotIn('Platform Price', headers)
        self.assertNotIn('Profit', headers)
        self.assertNotIn('Accumulated Line Revenue', headers)
        self.assertNotIn('Accumulated Profit', headers)
        labels = [row[4] for row in sheet.iter_rows(min_row=4, values_only=True)]
        self.assertEqual(labels, ['Sales Widget'])
        self.assertEqual(sheet['C4'].value, 'Statement Customer')
        self.assertEqual(sheet['D4'].value, 'STMT-W-1')
        self.assertEqual(sheet['F4'].value, 109.0)
        self.assertEqual(sheet['G4'].value, 1)
        self.assertEqual(sheet['H4'].value, '=F4*G4')
        self.assertEqual(sheet['I4'].value, '=IF($B$1=0,"",H4/$B$1)')

        monthly = workbook['Monthly Sales']
        self.assertEqual([cell.value for cell in monthly[1]], ['Year', 'Month', 'Revenue (USD)'])
        self.assertEqual(monthly['A2'].value, self.year)
        self.assertEqual(monthly['B2'].value, self.month)
        self.assertEqual(
            monthly['C2'].value,
            "=SUMIFS('Sale Transactions'!$I$4:$I$4,'Sale Transactions'!$B$4:$B$4,\">=\"&DATE(A2,B2,1),'Sale Transactions'!$B$4:$B$4,\"<\"&DATE(A2,B2+1,1))",
        )

        products = workbook['Product Sales']
        self.assertEqual(
            [cell.value for cell in products[1]],
            ['Product ID', 'Product Name', 'Year', 'Month', 'Quantity', 'Revenue (USD)'],
        )
        self.assertEqual(products['A2'].value, 'STMT-W-1')
        self.assertEqual(products['B2'].value, 'Sales Widget')
        self.assertEqual(products['C2'].value, self.year)
        self.assertEqual(products['D2'].value, self.month)
        self.assertEqual(
            products['E2'].value,
            "=SUMIFS('Sale Transactions'!$G$4:$G$4,'Sale Transactions'!$D$4:$D$4,A2,'Sale Transactions'!$E$4:$E$4,B2,'Sale Transactions'!$B$4:$B$4,\">=\"&DATE(C2,D2,1),'Sale Transactions'!$B$4:$B$4,\"<\"&DATE(C2,D2+1,1))",
        )
        self.assertEqual(
            products['F2'].value,
            "=SUMIFS('Sale Transactions'!$I$4:$I$4,'Sale Transactions'!$D$4:$D$4,A2,'Sale Transactions'!$E$4:$E$4,B2,'Sale Transactions'!$B$4:$B$4,\">=\"&DATE(C2,D2,1),'Sale Transactions'!$B$4:$B$4,\"<\"&DATE(C2,D2+1,1))",
        )

    def test_multiple_periods_include_only_selected_months(self):
        june = date(self.year - 1, 6, 15)
        july = date(self.year - 1, 7, 15)
        june_product = Product.objects.create(
            name='June Widget',
            sku='STMT-W-JUNE',
            reference_price=Decimal('12.00'),
        )
        july_product = Product.objects.create(
            name='July Widget',
            sku='STMT-W-JULY',
            reference_price=Decimal('13.00'),
        )
        for when, product in ((june, june_product), (july, july_product)):
            order = Order.objects.create(
                agent=self.agent,
                created_by=self.superuser,
                status=Order.OrderStatus.COMPLETED,
                customer_name='Statement Customer',
                transaction_date=when,
            )
            OrderItem.objects.create(
                order=order,
                product=product,
                quantity=1,
                selling_price=Decimal('20.00'),
                landed_cost=Decimal('5.00'),
            )
        selected = [
            f'{self.year}-{self.month:02d}',
            f'{june.year:04d}-{june.month:02d}',
        ]
        response = self.client.get('/order/export/', {'period': selected})
        self.assertEqual(response.status_code, 200, response.content[:200])
        sheet = load_workbook(BytesIO(response.content)).active
        labels = [row[4] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertIn('Sales Widget', labels)
        self.assertIn('June Widget', labels)
        self.assertNotIn('July Widget', labels)
        self.assertIn(
            f'order_statement_{june.year}_{june.month:02d}_{self.year}_{self.month:02d}.xlsx',
            response['Content-Disposition'],
        )

        simple = self.client.get('/order/export/', {
            'period': selected,
            'simple_report': '1',
            'usd_rate': '4.70',
        })
        self.assertEqual(simple.status_code, 200, simple.content[:200])
        monthly = load_workbook(BytesIO(simple.content))['Monthly Sales']
        self.assertEqual(
            [(monthly['A2'].value, monthly['B2'].value), (monthly['A3'].value, monthly['B3'].value)],
            [(june.year, june.month), (self.year, self.month)],
        )
        product_rows = [
            (row[0], row[2], row[3])
            for row in monthly.parent['Product Sales'].iter_rows(min_row=2, values_only=True)
        ]
        self.assertEqual(product_rows, [
            ('STMT-W-1', self.year, self.month),
            ('STMT-W-JUNE', june.year, june.month),
        ])

    def test_simple_report_requires_a_positive_rate(self):
        missing = self.client.get('/order/export/', {
            'month': self.month,
            'year': self.year,
            'simple_report': '1',
        })
        self.assertEqual(missing.status_code, 400)
        zero = self.client.get('/order/export/', {
            'month': self.month,
            'year': self.year,
            'simple_report': '1',
            'usd_rate': '0',
        })
        self.assertEqual(zero.status_code, 400)
